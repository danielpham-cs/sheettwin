"""Understanding features built on a language model, grounded in the SheetIR database.

The model never reads the file and never does arithmetic on its own. It reads a sketch of the parsed workbook
(sheetir.sketch) and either answers from it or asks for one read-only SQL query, which SheetIR runs on the database; the
answer is written from those rows. Cell ids in an answer are checked against the database. Table suggestions come back
in the format the user's own table corrections use, and are validated the same way before anyone sees them.
"""

from __future__ import annotations

import json
import re
import sqlite3

from openpyxl.utils import get_column_letter, range_boundaries

from .llm import LLM, LLMError
from .sketch import sketch

PROMPT_VERSION = 1
MAX_SQL_ROWS = 60
_CELL = re.compile(r"(?:'[^']+'|[A-Za-z_][\w.]*)![A-Z]{1,3}\d+")

SYSTEM = """You help people understand an Excel workbook that SheetIR has parsed into a SQLite database.
You get a sketch of the workbook (sheets, tables with their columns and formula patterns, parameters, results,
lookups, issues) and the database schema. Ground every statement in that material: never invent cells, values or
formulas. Cite cells by their ids exactly as written (for example Quote!J13; a sheet name with spaces or symbols is
quoted: 'Class Data'!A4). Explain what the workbook does in
business terms when asked. Formula patterns are in R1C1 notation: RC[-1] is the cell one column to the left in the
same row. Reply in the language of the user's question."""

SCHEMA_NOTES = """Database (SQLite, read-only). Query cell contents through cell_values:
cell_values(cell_id, sheet, row, col, a1, vtype, role, val, formula, semantic, table_name, field, part, group_id)
  val is the plain value (text or number; for a formula, the result Excel last stored, NULL if none stored);
  table_name and field are the names shown in the sketch (e.g. table_name = 'quote' AND field = 'margin');
  part is 'header', 'body' or 'total' inside a table, NULL outside; role is input|derived|output|label|orphan;
  semantic is a readable address such as quote["Cover plate"].margin.
  Example, a column of a table with each row's key:
    SELECT k.val AS key, v.val FROM cell_values v JOIN cell_values k ON k.table_name = v.table_name AND k.row = v.row
      AND k.col = (SELECT key_col FROM tables WHERE name = v.table_name)
    WHERE v.table_name = 'quote' AND v.field = 'margin' AND v.part = 'body'
tables(table_id, sheet, name, kind, source, header_row, row_start, row_end, col_start, col_end, key_col, total_row, ...)
fields(table_id, col, header, name, unit, kind, r1c1)       params(cell_id, label, name, unit, domain)
formula_groups(group_id, sheet, axis, anchor, last, n_cells, r1c1)
issues(cell_id, kind, severity, detail)                      relations(...), table_links (view)
deps(src_cell, dst_cell)  (view: dst_cell's formula reads src_cell)"""


def _rows(con: sqlite3.Connection, sql: str) -> tuple[list[str], list[list]]:
    q = sql.strip().rstrip(";")
    if not re.match(r"^(select|with)\b", q, re.I):
        raise ValueError("only SELECT / WITH queries are allowed")
    cur = con.execute(q)
    return [d[0] for d in cur.description or []], [list(r) for r in cur.fetchmany(MAX_SQL_ROWS)]


def _check_cells(con: sqlite3.Connection, text: str) -> tuple[list[str], list[str]]:
    """Cell ids mentioned in a text: those in the database, and those that are not."""
    found = list(dict.fromkeys(_CELL.findall(text)))
    known = {r[0] for r in con.execute(f"SELECT cell_id FROM cells WHERE cell_id IN ({','.join('?' * len(found))})", found)} if found else set()
    return [c for c in found if c in known], [c for c in found if c not in known]


def _table_views(con: sqlite3.Connection) -> str:
    """One temp view per table, a row per body row and a column per field (plain values), so questions about rows
    read as ordinary SQL: SELECT student_name FROM table_class_data WHERE gender = 'Female'. Returns their description."""
    lines = []
    for tid, name, total in con.execute("SELECT table_id, name, total_row FROM tables ORDER BY table_id").fetchall():
        fields = con.execute("SELECT col, name FROM fields WHERE table_id = ? ORDER BY col", (tid,)).fetchall()
        if not fields:
            continue
        view = "table_" + re.sub(r"\W", "_", name)
        cols = ", ".join(f"MAX(CASE WHEN col = {c} THEN val END) AS \"{f}\"" for c, f in fields)
        con.execute(f'CREATE TEMP VIEW IF NOT EXISTS "{view}" AS SELECT row, {cols} FROM cell_values '
                    f"WHERE table_name = '{name.replace(chr(39), chr(39) * 2)}' AND part = 'body' GROUP BY row")
        lines.append(f"{view}(row, {', '.join(f for _, f in fields)})" + (f"  -- total row {total} excluded" if total else ""))
    if not lines:
        return ""
    return ("\n\n# Table views (prefer these for questions about rows: one row per table row, plain values)\n"
            + "\n".join(lines))


PLAN_SCHEMA = {
    "type": "object", "additionalProperties": False, "required": ["answer", "sql"],
    "properties": {
        "answer": {"type": "string", "description": "markdown answer if the sketch is enough, else empty"},
        "sql": {"type": "string", "description": "one read-only SQLite query if data is needed, else empty"},
    },
}
ANSWER_SCHEMA = {
    "type": "object", "additionalProperties": False, "required": ["answer"],
    "properties": {"answer": {"type": "string", "description": "markdown answer"}},
}


NO_RESULTS = {
    "en": "Note: this file stores no formula results (it was not saved by Excel), so questions about computed columns "
          "cannot be answered from it. Open the file in Excel, save it, and upload it again.",
    "zh": "注意：這個檔案沒有儲存任何公式的計算結果（不是由 Excel 存檔的），所以和計算欄位有關的問題無法從檔案回答。"
          "請用 Excel 開啟並存檔後再上傳。",
}


def ask(con: sqlite3.Connection, question: str, llm: LLM, lang: str = "en") -> dict:
    """Answer a question about the workbook. Returns {reply, cells, sql, unknown_cells, model}."""
    context = f"{SCHEMA_NOTES}\n\n# Sketch\n{sketch(con)}"
    con.row_factory = None
    con.execute("CREATE TEMP VIEW IF NOT EXISTS cell_values AS SELECT c.cell_id, c.sheet, c.row, c.col, c.a1, c.vtype, "
                "c.role, CASE WHEN c.formula IS NULL THEN json_extract(c.value, '$') ELSE json_extract(c.cached, '$') END AS val, "
                "c.formula, c.semantic, t.name AS table_name, f.name AS field, "
                "CASE WHEN t.table_id IS NULL THEN NULL WHEN c.row <= COALESCE(t.header_row, 0) THEN 'header' "
                "WHEN c.row = t.total_row THEN 'total' ELSE 'body' END AS part, c.group_id "
                "FROM cells c LEFT JOIN tables t ON t.table_id = c.table_id "
                "LEFT JOIN fields f ON f.table_id = c.table_id AND f.col = c.col")
    context += _table_views(con)
    ask_q = f"{context}\n\n# Question\n{question}"
    plan = llm.json(SYSTEM + "\nAnswer directly in `answer` (leave `sql` empty) only questions about what the workbook "
                    "is, its structure, its formulas and its issues. The sketch shows only a few sample rows of each "
                    "table, so any question about particular rows, counts, totals, lists or values needs data: then "
                    "leave `answer` empty and write one SQLite SELECT in `sql` (LIMIT 60 at most) that returns "
                    "everything the answer needs, including the rows themselves when the question asks who or which. "
                    "For the most or least of something, keep ties (compare with a MAX/MIN subquery rather than LIMIT 1). "
                    "Never tell the user that a query is needed: write it.", ask_q, PLAN_SCHEMA, "plan")
    sql, answer, columns, rows = plan["sql"].strip(), plan["answer"].strip(), [], []
    if sql:
        # one revision when the query fails or finds nothing: an empty result is more often a wrong query than a fact
        for attempt in range(2):
            try:
                columns, rows = _rows(con, sql)
                problem = None if rows or attempt else "returned no rows"
            except (sqlite3.Error, ValueError) as e:
                columns, rows, problem = [], [], f"failed: {e}"
            if not problem:
                break
            fix = llm.json(SYSTEM, f"{ask_q}\n\nYour query\n{sql}\n{problem}. Check table, column and value names "
                           "against the sketch (cell_values.val holds plain values) and write a corrected SQLite SELECT "
                           "in `sql`; leave `answer` empty.", PLAN_SCHEMA, "plan")
            if not fix["sql"].strip():
                break
            sql = fix["sql"].strip()
        data = json.dumps({"columns": columns, "rows": rows}, ensure_ascii=False, default=str)
        answer = llm.json(SYSTEM + "\nWrite the answer from the query result; quote numbers exactly as they appear "
                          "there. If the result is empty, say that no matching data was found, not that none exists.",
                          f"{ask_q}\n\n# Query\n{sql}\n\n# Result (JSON)\n{data}", ANSWER_SCHEMA, "answer")["answer"].strip()
        # Said by SheetIR, not left to the model: an empty or NULL answer may only mean Excel never stored the results.
        n_f, n_cached = con.execute("SELECT COUNT(*), COUNT(cached) FROM cells WHERE vtype = 'formula'").fetchone()
        if n_f and not n_cached and (not rows or any(v is None for r in rows for v in r)):
            answer += "\n\n> " + NO_RESULTS.get(lang, NO_RESULTS["en"])
    cells, unknown = _check_cells(con, answer)
    return {"reply": answer, "cells": cells, "unknown_cells": unknown, "sql": sql or None, "model": llm.model}


# ---- table suggestions ---------------------------------------------------------------------------

GRID_HEAD, GRID_TAIL, GRID_COLS = 50, 8, 26

SUGGEST_SCHEMA = {
    "type": "object", "additionalProperties": False, "required": ["tables"],
    "properties": {"tables": {"type": "array", "items": {
        "type": "object", "additionalProperties": False,
        "required": ["range", "header_rows", "total", "name", "reason"],
        "properties": {
            "range": {"type": "string", "description": "A1 range from the first header row to the last row (incl. total)"},
            "header_rows": {"type": "integer", "description": "how many of the first rows are header rows (0-5)"},
            "total": {"type": "boolean", "description": "whether the last row is a total row"},
            "name": {"type": "string", "description": "short snake_case name for the table"},
            "reason": {"type": "string", "description": "one sentence"},
        }}}},
}


def _grid(con: sqlite3.Connection, sheet: str) -> tuple[str, int, int]:
    """The sheet as rows of cells: "B3:'Copper price'" (text), "B4:285" (number), "C5:ƒ=SUM(…)" (formula)."""
    cells = con.execute("SELECT row, col, vtype, value, formula FROM cells WHERE sheet = ? ORDER BY row, col", (sheet,)).fetchall()
    max_r = max((c[0] for c in cells), default=0)
    max_c = max((c[1] for c in cells), default=0)
    by_row: dict[int, list] = {}
    for r, c, vt, v, f in cells:
        if c > GRID_COLS:
            continue
        a1 = f"{get_column_letter(c)}{r}"
        if f:
            txt = "ƒ" + f[:30]
        else:
            val = json.loads(v) if v else ""
            txt = repr(str(val)[:30]) if vt == "text" else str(val)[:20]
        by_row.setdefault(r, []).append(f"{a1}:{txt}")
    rows = sorted(by_row)
    shown = rows if len(rows) <= GRID_HEAD + GRID_TAIL else rows[:GRID_HEAD] + [None] + rows[-GRID_TAIL:]
    lines = [f"(rows {rows[GRID_HEAD]}-{rows[-GRID_TAIL - 1]} not shown; they continue the pattern)" if r is None
             else f"{r}: " + "  ".join(by_row[r]) for r in shown]
    if max_c > GRID_COLS:
        lines.append(f"(columns after {get_column_letter(GRID_COLS)} not shown; the sheet goes to column {get_column_letter(max_c)})")
    return "\n".join(lines), max_r, max_c


def suggest_tables(con: sqlite3.Connection, sheet: str, llm: LLM) -> dict:
    """Proposed tables for a sheet, as user-table specs; invalid ones are dropped with the reason."""
    grid, max_r, max_c = _grid(con, sheet)
    current = [f"{t[0]} {get_column_letter(t[3])}{t[1]}:{get_column_letter(t[4])}{t[2]}" for t in con.execute(
        "SELECT name, COALESCE(header_top, header_row, row_start), COALESCE(total_row, row_end), col_start, col_end "
        "FROM tables WHERE sheet = ?", (sheet,))]
    prompt = (f"Sheet {sheet!r}, up to row {max_r}, column {get_column_letter(max(max_c, 1))}. Non-empty cells by row "
              f"(text in quotes, ƒ = formula):\n{grid}\n\nTables detected by rules (may be wrong or missing): "
              f"{', '.join(current) or 'none'}\n\nFind every data table on this sheet. A table is a rectangular block "
              "with header rows on top (column labels, possibly several levels, possibly year labels) and data rows "
              "below; row labels in the first column(s) belong to the table even when their header cell is empty. "
              "Titles, notes, sources and footnotes above or below are not part of a table. Blank rows inside a table "
              "between sections do not end it. Return tables that do not overlap.")
    res = llm.json(SYSTEM, prompt, SUGGEST_SCHEMA, "tables")
    valid, dropped, boxes = [], [], []
    for t in res["tables"]:
        try:
            c1, r1, c2, r2 = range_boundaries(t["range"].replace("$", "").upper())
            if None in (c1, r1, c2, r2):
                raise ValueError
        except ValueError:
            dropped.append({**t, "why": "not a cell range"})
            continue
        n = max(0, min(5, t["header_rows"]))
        if (r2 - r1 + 1) - n - (1 if t["total"] else 0) < 1:
            dropped.append({**t, "why": "no body row left"})
            continue
        if any(c1 <= d2 and d1 <= c2 and r1 <= s2 and s1 <= r2 for d1, s1, d2, s2 in boxes):
            dropped.append({**t, "why": "overlaps another suggestion"})
            continue
        boxes.append((c1, r1, c2, r2))
        valid.append({"range": f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}", "header_rows": n,
                      "total": bool(t["total"]), "name": t["name"] or None, "reason": t["reason"]})
    return {"tables": valid, "dropped": dropped, "model": llm.model}


__all__ = ["ask", "suggest_tables", "LLMError"]
