"""L0: read the raw workbook (formulas and cached values) without modifying it."""

from __future__ import annotations

import datetime as dt
import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

CellKey = tuple[str, int, int]  # (sheet, row, col)


@dataclass
class RawCell:
    sheet: str
    row: int
    col: int
    vtype: str  # formula | number | text | bool | date | error
    value: object  # constant value; None for formulas
    formula: str | None
    cached: object  # value Excel last computed (None if never recalculated)
    number_format: str
    flags: list[str] = field(default_factory=list)
    style: frozenset[str] = frozenset()  # header-like formatting: bold | fill | border_bottom


@dataclass
class TableDef:
    """An Excel Table (ListObject) declared in the file."""

    name: str
    sheet: str
    ref: str
    header_rows: int
    totals_rows: int
    columns: list[str]


@dataclass
class DefinedName:
    name: str
    scope: str | None  # sheet name for sheet-scoped names
    refers_to: str


@dataclass
class Validation:
    sheet: str
    sqref: str
    rule: dict


@dataclass
class RawWorkbook:
    path: Path
    sha256: str
    sheets: list[str]
    cells: dict[CellKey, RawCell]
    names: list[DefinedName]
    tables: list[TableDef]
    validations: list[Validation]
    merged: list[tuple[str, str]]


def _vtype(cell) -> str:
    if cell.data_type == "f":
        return "formula"
    if cell.data_type == "e":
        return "error"
    if cell.data_type == "b" or isinstance(cell.value, bool):
        return "bool"
    if cell.data_type == "d" or isinstance(cell.value, (dt.date, dt.datetime, dt.time)):
        return "date"
    if cell.data_type == "n":
        return "number"
    return "text"


def _style(cell) -> frozenset[str]:
    if not cell.has_style:
        return frozenset()
    marks = set()
    if cell.font is not None and cell.font.b:
        marks.add("bold")
    if cell.fill is not None and cell.fill.fill_type not in (None, "none"):
        marks.add("fill")
    if cell.border is not None and cell.border.bottom is not None and cell.border.bottom.style:
        marks.add("border_bottom")
    return frozenset(marks)


def load(path: str | Path) -> RawWorkbook:
    path = Path(path)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)

    cells: dict[CellKey, RawCell] = {}
    tables: list[TableDef] = []
    validations: list[Validation] = []
    merged: list[tuple[str, str]] = []
    names: list[DefinedName] = [DefinedName(n, None, d.attr_text) for n, d in wb_f.defined_names.items()]

    for ws in wb_f.worksheets:
        ws_v = wb_v[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                flags: list[str] = []
                value = cell.value
                formula = None
                if isinstance(value, ArrayFormula):
                    formula, flags = value.text, ["array_formula"]
                elif isinstance(value, DataTableFormula):
                    formula, flags = "=" + str(value), ["data_table"]
                elif cell.data_type == "f":
                    formula = value
                vtype = "formula" if formula is not None else _vtype(cell)
                cells[(ws.title, cell.row, cell.column)] = RawCell(
                    sheet=ws.title,
                    row=cell.row,
                    col=cell.column,
                    vtype=vtype,
                    value=None if formula is not None else value,
                    formula=formula,
                    cached=ws_v.cell(cell.row, cell.column).value if formula is not None else value,
                    number_format=cell.number_format,
                    flags=flags,
                    style=_style(cell),
                )

        for t in ws.tables.values():
            tables.append(TableDef(
                name=t.displayName or t.name,
                sheet=ws.title,
                ref=t.ref,
                header_rows=1 if t.headerRowCount is None else t.headerRowCount,
                totals_rows=t.totalsRowCount or 0,
                columns=[c.name for c in t.tableColumns],
            ))
        for dv in ws.data_validations.dataValidation:
            rule = {"type": dv.type, "operator": dv.operator, "formula1": dv.formula1, "formula2": dv.formula2}
            validations.append(Validation(ws.title, str(dv.sqref), rule))
        merged.extend((ws.title, str(r)) for r in ws.merged_cells.ranges)
        names.extend(DefinedName(n, ws.title, d.attr_text) for n, d in ws.defined_names.items())

    return RawWorkbook(path, sha, wb_f.sheetnames, cells, names, tables, validations, merged)
